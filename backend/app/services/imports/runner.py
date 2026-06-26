"""Import wizard orchestration service (S10).

Public API
----------
translate_column_map(wizard_map)          -> dict      (S06-compatible column_map)
run_preview(db, batch, ...)               -> counts dict
run_import(db, batch, ...)                -> counts dict + elapsed_ms
stream_import_report_csv(db, batch_id)   -> AsyncIterator[str]

Design notes
------------
- translate_column_map converts wizard format (custom:<name>, ignore) to S06
  format (custom_data.<name>, absent) so that S06 mapper functions can be reused
  without modification.
- run_preview never writes core table rows.  ImportRowResult rows are added
  AFTER each savepoint resolves (S06 dry_run pattern, CN-13).
- run_import commits in chunks of IMPORT_CHUNK_SIZE rows (default 500).
- stream_import_report_csv applies _sanitize_cell to text cells to defuse
  formula injection (AC-report-sanitise).
- After a successful run_import, member_status_service.recompute_all_contacts
  is called (CN-24) guarded with ImportError.
- Circular-import rule (CN-25): must NOT import from any router module.
"""
from __future__ import annotations

import asyncio
import csv
import io
import logging
import time
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, AsyncIterator, Optional

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.config import legacy_settings
from app.models import Contact, ImportBatch, ImportRowResult, utc_now
from app.services.migration.mapper import (
    MappingError,
    map_contact_row,
    map_participant_row,
)
from app.services.migration.loader import upsert_contact
from app.services.imports.staging import iter_csv_rows
from app.services.imports.disposition import classify

logger = logging.getLogger(__name__)

IMPORT_CHUNK_SIZE: int = int(
    getattr(legacy_settings, "IMPORT_CHUNK_SIZE", 500)
    if hasattr(legacy_settings, "IMPORT_CHUNK_SIZE")
    else 500
)

_ENTITY_SINGULAR: dict[str, str] = {
    "contacts":     "contact",
    "events":       "event",
    "participants": "participant",
    "links":        "contact",
}


# ---------------------------------------------------------------------------
# Column-map translation
# ---------------------------------------------------------------------------


def translate_column_map(wizard_map: dict[str, Any]) -> dict[str, Any]:
    """Convert the wizard column_map format to the S06 mapper format.

    Transformations
    ---------------
    - ``custom:<name>``  →  ``custom_data.<name>``  (rewrite target)
    - ``ignore``         →  entry dropped entirely
    - all other targets  →  unchanged
    """
    s06_map: dict[str, Any] = {}
    for header, spec in wizard_map.items():
        if not isinstance(spec, dict):
            continue
        target = spec.get("target", "ignore")
        if target == "ignore":
            continue
        if target.startswith("custom:"):
            field_name = target[len("custom:"):]
            new_spec = dict(spec)
            new_spec["target"] = f"custom_data.{field_name}"
            s06_map[header] = new_spec
        else:
            s06_map[header] = spec
    return s06_map


# ---------------------------------------------------------------------------
# XLSX row helper (sheet-aware, mirrors reader._iter_rows_sync)
# ---------------------------------------------------------------------------


def _iter_xlsx_sync(path: str, sheet: Optional[str] = None) -> list[dict]:
    """Read XLSX rows from *sheet* (or active sheet) synchronously."""
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet_names = wb.sheetnames
        if sheet and sheet in sheet_names:
            ws = wb[sheet]
        else:
            ws = wb.active

        rows_iter = ws.iter_rows(values_only=True)
        try:
            raw_headers = next(rows_iter)
        except StopIteration:
            return []

        # Build (position, name) pairs for non-empty headers
        header_positions: list[tuple[int, str]] = []
        for i, cell in enumerate(raw_headers):
            val = str(cell).strip() if cell is not None else ""
            if val:
                header_positions.append((i, val))

        results: list[dict] = []
        for raw_row in rows_iter:
            row_dict: dict[str, Any] = {}
            all_blank = True
            for col_pos, col_name in header_positions:
                cell_val = raw_row[col_pos] if col_pos < len(raw_row) else None
                if isinstance(cell_val, str):
                    cell_val = cell_val.strip() or None
                if cell_val is not None:
                    all_blank = False
                row_dict[col_name] = cell_val
            if all_blank:
                continue
            results.append(row_dict)
        return results
    finally:
        wb.close()


# ---------------------------------------------------------------------------
# run_preview
# ---------------------------------------------------------------------------


async def run_preview(
    db: AsyncSession,
    batch: ImportBatch,
    column_map: dict[str, Any],
    match_key: str,
    conflict_policy: str,
    sheet: Optional[str] = None,
    target_event_id: Optional[int] = None,
) -> dict[str, int]:
    """Simulate an import run without writing any core table rows.

    Writes ImportRowResult rows so the UI can inspect per-row outcomes.
    Returns counts: would_create, would_update, would_skip, ambiguous, errors.
    """
    s06_map = translate_column_map(column_map)
    entity = _ENTITY_SINGULAR.get(batch.entity, batch.entity)
    staging_abs = Path(legacy_settings.STORAGE_PATH) / batch.staging_file
    fmt = (batch.options or {}).get("fmt", "xlsx")

    batch.mode = "wizard_preview"
    batch.status = "running"
    await db.flush()

    counts: dict[str, int] = {
        "would_create": 0,
        "would_update": 0,
        "would_skip":   0,
        "ambiguous":    0,
        "errors":       0,
    }
    row_number = 0
    buffer: list[tuple[int, dict, Any]] = []

    async def _flush(buf: list[tuple[int, dict, Any]]) -> None:
        if not buf:
            return
        pending: list[ImportRowResult] = []
        # Savepoint wraps classification SELECTs; rolled back to discard any
        # accidental writes and release sub-transaction resources.
        sp = await db.begin_nested()
        try:
            for rn, raw_row, mapped_or_err in buf:
                if _is_mapping_error(mapped_or_err):
                    counts["errors"] += 1
                    pending.append(_make_irr(
                        batch.id, rn,
                        str(raw_row.get("external_id", "")),
                        "error", None,
                        getattr(mapped_or_err, "message", str(mapped_or_err)),
                    ))
                    continue

                if entity == "contact":
                    kind, existing, msg = await classify(mapped_or_err, match_key, db)
                    outcome, eid = _preview_outcome(kind, existing, conflict_policy, counts)
                    pending.append(_make_irr(
                        batch.id, rn,
                        str(getattr(mapped_or_err, "external_id", "") or ""),
                        outcome, eid, msg,
                    ))
                else:
                    # participants / events: count all valid mappings as would_create
                    counts["would_create"] += 1
                    ext_raw = raw_row.get("contact_ref") or raw_row.get("external_id", "")
                    pending.append(_make_irr(
                        batch.id, rn, str(ext_raw),
                        "would_create", None, "preview",
                    ))

            await sp.rollback()
        except Exception:
            try:
                await sp.rollback()
            except Exception:
                pass
            raise

        # Add audit rows AFTER savepoint resolves so a rollback does not discard them
        for irr in pending:
            db.add(irr)

    try:
        async for raw_row in _row_generator(fmt, staging_abs, sheet, batch.options):
            row_number += 1
            mapped = await _map_row(entity, raw_row, s06_map)
            buffer.append((row_number, raw_row, mapped))
            if len(buffer) >= IMPORT_CHUNK_SIZE:
                await _flush(buffer)
                await db.commit()
                buffer = []
    except ValueError as exc:
        # Mid-stream decode error: flush what we have, record batch failure.
        logger.warning("run_preview row-iteration error at row %s: %s", row_number, exc)
        counts["errors"] += 1
        if buffer:
            await _flush(buffer)
            await db.commit()
            buffer = []
        irr = _make_irr(batch.id, row_number + 1, "", "error", None, str(exc)[:1000])
        db.add(irr)

    if buffer:
        await _flush(buffer)
        await db.commit()

    batch.total_rows = row_number
    batch.created_count = counts["would_create"]
    batch.updated_count = counts["would_update"]
    batch.skipped_count = counts["would_skip"]
    batch.error_count = counts["errors"]
    batch.review_count = counts["ambiguous"]
    batch.status = "completed"
    batch.finished_at = utc_now()
    await db.commit()

    return counts


# ---------------------------------------------------------------------------
# run_import
# ---------------------------------------------------------------------------


async def run_import(
    db: AsyncSession,
    batch: ImportBatch,
    column_map: dict[str, Any],
    match_key: str,
    conflict_policy: str,
    sheet: Optional[str] = None,
    target_event_id: Optional[int] = None,
) -> dict[str, Any]:
    """Execute the actual import, writing core table rows.

    Returns counts dict plus elapsed_ms and status.
    Idempotent under conflict_policy='skip' (re-run adds 0 net rows).
    After completion calls member_status_service.recompute_all_contacts (CN-24).
    """
    from app.services.bulk_service import bulk_upsert_participants

    s06_map = translate_column_map(column_map)
    entity = _ENTITY_SINGULAR.get(batch.entity, batch.entity)
    staging_abs = Path(legacy_settings.STORAGE_PATH) / batch.staging_file
    fmt = (batch.options or {}).get("fmt", "xlsx")

    batch.mode = "wizard_run"
    batch.status = "running"
    await db.flush()

    counts: dict[str, int] = {
        "created": 0,
        "updated": 0,
        "skipped": 0,
        "review":  0,
        "errors":  0,
    }
    row_number = 0
    buffer: list[tuple[int, dict, Any]] = []
    # For participant rows we accumulate valid dicts; errors go straight to pending
    participant_buffer: list[tuple[int, dict, Any]] = []

    t_start = time.monotonic()

    # For participant entity with a target_event_id: inject a synthetic event_ref
    # column so that map_participant_row (which requires event_ref in the row) does
    # not return a MappingError when the user relies on the API-level event selector.
    _synthetic_event_key: Optional[str] = None
    if entity == "participant" and target_event_id is not None:
        has_event_ref_mapped = any(
            isinstance(spec, dict) and spec.get("target") == "event_ref"
            for spec in s06_map.values()
        )
        if not has_event_ref_mapped:
            _synthetic_event_key = "_wizard_event_ref_"
            s06_map = dict(s06_map)
            s06_map[_synthetic_event_key] = {"target": "event_ref", "data_type": "text"}

    async def _flush_contacts(buf: list[tuple[int, dict, Any]]) -> None:
        if not buf:
            return
        pending: list[ImportRowResult] = []
        for rn, raw_row, mapped_or_err in buf:
            if _is_mapping_error(mapped_or_err):
                counts["errors"] += 1
                pending.append(_make_irr(
                    batch.id, rn,
                    str(raw_row.get("external_id", "")),
                    "error", None,
                    getattr(mapped_or_err, "message", str(mapped_or_err)),
                ))
                continue

            try:
                kind, existing, disp_msg = await classify(mapped_or_err, match_key, db)
                if kind == "error":
                    counts["errors"] += 1
                    pending.append(_make_irr(
                        batch.id, rn,
                        str(getattr(mapped_or_err, "external_id", "") or ""),
                        "error", None, disp_msg,
                    ))

                elif kind == "new":
                    if match_key == "external_id":
                        # Normal path: upsert keyed on external_id.
                        outcome, eid, msg = await upsert_contact(db, dict(vars(mapped_or_err)))
                    else:
                        # Email/name match key: disposition said no existing contact
                        # with that email/name, but the row may still carry an
                        # external_id that matches a different contact.  Using
                        # upsert_contact would silently UPDATE that contact (wrong).
                        # Instead always INSERT a fresh row.
                        outcome, eid, msg = await _create_contact_direct(db, mapped_or_err)
                    _tally_import(counts, outcome)
                    pending.append(_make_irr(
                        batch.id, rn,
                        str(getattr(mapped_or_err, "external_id", "") or ""),
                        outcome, eid, msg,
                    ))

                elif kind == "match":
                    assert existing is not None
                    outcome, msg = await _apply_match(db, existing, mapped_or_err, conflict_policy)
                    _tally_import(counts, outcome)
                    pending.append(_make_irr(
                        batch.id, rn,
                        str(getattr(mapped_or_err, "external_id", "") or ""),
                        outcome, existing.id, msg,
                    ))

                elif kind == "ambiguous":
                    # Enqueue for manual review
                    counts["review"] += 1
                    raw_name = (
                        f"{getattr(mapped_or_err, 'first_name', '') or ''} "
                        f"{getattr(mapped_or_err, 'last_name', '') or ''}"
                    ).strip()
                    review_id = await _enqueue_name_review(raw_name, batch.id, db)
                    pending.append(_make_irr(
                        batch.id, rn,
                        str(getattr(mapped_or_err, "external_id", "") or ""),
                        "review", None,
                        f"ambiguous name match for {raw_name!r}; review_queue_id={review_id}",
                    ))

            except Exception as exc:
                counts["errors"] += 1
                logger.exception(
                    "run_import: row %s error: %s", rn, exc
                )
                pending.append(_make_irr(
                    batch.id, rn,
                    str(raw_row.get("external_id", "")),
                    "error", None, str(exc)[:500],
                ))

        await db.flush()
        for irr in pending:
            db.add(irr)

    async def _flush_participants(buf: list[tuple[int, dict, Any]]) -> None:
        if not buf:
            return
        pending: list[ImportRowResult] = []
        valid_rows: list[dict[str, Any]] = []
        valid_meta: list[tuple[int, dict]] = []

        for rn, raw_row, mapped_or_err in buf:
            if _is_mapping_error(mapped_or_err):
                counts["errors"] += 1
                pending.append(_make_irr(
                    batch.id, rn,
                    str(raw_row.get("contact_ref", "")),
                    "error", None,
                    getattr(mapped_or_err, "message", str(mapped_or_err)),
                ))
                continue

            # Resolve contact_ref to contact_id
            contact_ref = str(getattr(mapped_or_err, "contact_ref", "") or "")
            event_ref = str(getattr(mapped_or_err, "event_ref", "") or "")

            contact_id = await _resolve_contact_ref(contact_ref, db)
            if contact_id is None:
                counts["errors"] += 1
                pending.append(_make_irr(
                    batch.id, rn, contact_ref,
                    "error", None,
                    f"contact_ref {contact_ref!r} not found",
                ))
                continue

            event_id = target_event_id
            if event_id is None:
                event_id = await _resolve_event_ref(event_ref, db)
            if event_id is None:
                counts["errors"] += 1
                pending.append(_make_irr(
                    batch.id, rn, contact_ref,
                    "error", None,
                    f"event_ref {event_ref!r} not found and no target_event_id",
                ))
                continue

            valid_rows.append({
                "contact_id": contact_id,
                "event_id":   event_id,
                "status":     getattr(mapped_or_err, "status", "attended"),
                "role":       getattr(mapped_or_err, "role", None),
                "source":     "import",
            })
            valid_meta.append((rn, raw_row))

        if valid_rows:
            try:
                await bulk_upsert_participants(db, valid_rows)
                counts["created"] += len(valid_rows)
                for (rn, raw_row) in valid_meta:
                    pending.append(_make_irr(
                        batch.id, rn,
                        str(raw_row.get("contact_ref", "")),
                        "created", None, "participant upserted",
                    ))
            except Exception as exc:
                counts["errors"] += len(valid_rows)
                logger.exception("run_import: bulk_upsert_participants failed: %s", exc)
                for (rn, raw_row) in valid_meta:
                    pending.append(_make_irr(
                        batch.id, rn,
                        str(raw_row.get("contact_ref", "")),
                        "error", None, str(exc)[:500],
                    ))

        await db.flush()
        for irr in pending:
            db.add(irr)

    try:
        async for raw_row in _row_generator(fmt, staging_abs, sheet, batch.options):
            row_number += 1
            # Inject synthetic event_ref when needed (participant + target_event_id)
            effective_row = raw_row
            if _synthetic_event_key and target_event_id is not None:
                effective_row = {**raw_row, _synthetic_event_key: str(target_event_id)}
            mapped = await _map_row(entity, effective_row, s06_map)

            if entity == "contact":
                buffer.append((row_number, raw_row, mapped))
                if len(buffer) >= IMPORT_CHUNK_SIZE:
                    await _flush_contacts(buffer)
                    await db.commit()
                    buffer = []
            elif entity == "participant":
                participant_buffer.append((row_number, effective_row, mapped))
                if len(participant_buffer) >= IMPORT_CHUNK_SIZE:
                    await _flush_participants(participant_buffer)
                    await db.commit()
                    participant_buffer = []
            else:
                # Unsupported entity — record error
                counts["errors"] += 1
                irr = _make_irr(
                    batch.id, row_number, str(raw_row.get("external_id", "")),
                    "error", None, f"Unsupported entity: {entity!r}",
                )
                db.add(irr)
                if row_number % IMPORT_CHUNK_SIZE == 0:
                    await db.commit()
    except ValueError as exc:
        # Mid-stream decode error: flush pending rows and record failure.
        logger.warning("run_import row-iteration error at row %s: %s", row_number, exc)
        counts["errors"] += 1
        if entity == "contact" and buffer:
            await _flush_contacts(buffer)
            await db.commit()
            buffer = []
        elif entity == "participant" and participant_buffer:
            await _flush_participants(participant_buffer)
            await db.commit()
            participant_buffer = []
        irr = _make_irr(batch.id, row_number + 1, "", "error", None, str(exc)[:1000])
        db.add(irr)
        await db.commit()

    # Flush remaining rows
    if entity == "contact" and buffer:
        await _flush_contacts(buffer)
        await db.commit()
    elif entity == "participant" and participant_buffer:
        await _flush_participants(participant_buffer)
        await db.commit()

    elapsed_ms = int((time.monotonic() - t_start) * 1000)

    batch.total_rows = row_number
    batch.created_count = counts["created"]
    batch.updated_count = counts["updated"]
    batch.skipped_count = counts["skipped"]
    batch.error_count = counts["errors"]
    batch.review_count = counts["review"]
    batch.status = "completed"
    batch.finished_at = utc_now()
    await db.commit()

    # CN-24: recompute member status after live run
    try:
        from app.services import member_status_service  # noqa: PLC0415
        await member_status_service.recompute_all_contacts(db)
    except ImportError:
        pass

    return {**counts, "elapsed_ms": elapsed_ms, "status": "completed"}


# ---------------------------------------------------------------------------
# stream_import_report_csv
# ---------------------------------------------------------------------------


async def stream_import_report_csv(
    db: AsyncSession,
    batch_id: int,
) -> AsyncIterator[str]:
    """Yield a CSV header then one sanitised line per ImportRowResult.

    Uses db.stream + partitions(1000) for bounded memory.
    Applies _sanitize_cell to text fields (formula-injection guard).
    """
    from app.services.export_service import _sanitize_cell

    _REPORT_HEADERS = ["row_number", "external_id", "outcome", "entity_id", "message"]

    def _row_line(row: ImportRowResult) -> str:
        buf = io.StringIO()
        writer = csv.writer(buf, lineterminator="\n")
        writer.writerow([
            row.row_number,
            _sanitize_cell(row.external_id or ""),
            _sanitize_cell(row.outcome or ""),
            row.entity_id if row.entity_id is not None else "",
            _sanitize_cell(row.message or ""),
        ])
        return buf.getvalue()

    # Header
    buf = io.StringIO()
    csv.writer(buf, lineterminator="\n").writerow(_REPORT_HEADERS)
    yield buf.getvalue()

    stmt = (
        select(ImportRowResult)
        .where(ImportRowResult.batch_id == batch_id)
        .order_by(ImportRowResult.row_number)
        .execution_options(yield_per=1000)
    )
    result = await db.stream(stmt)
    async for partition in result.scalars().partitions(1000):
        for row in partition:
            yield _row_line(row)


# ---------------------------------------------------------------------------
# Internal helpers
# ---------------------------------------------------------------------------


def _is_mapping_error(obj: Any) -> bool:
    return hasattr(obj, "message") and isinstance(obj, MappingError)


def _make_irr(
    batch_id: int,
    row_number: int,
    external_id: str,
    outcome: str,
    entity_id: Optional[int],
    message: str,
) -> ImportRowResult:
    return ImportRowResult(
        batch_id=batch_id,
        row_number=row_number,
        external_id=external_id[:64] if external_id else "",
        outcome=outcome,
        entity_id=entity_id,
        message=message[:1000] if message else None,
        raw=None,
    )


def _preview_outcome(
    kind: str,
    existing: Any,
    conflict_policy: str,
    counts: dict[str, int],
) -> tuple[str, Optional[int]]:
    if kind == "error":
        counts["errors"] += 1
        return "error", None
    if kind == "new":
        counts["would_create"] += 1
        return "would_create", None
    if kind == "match":
        eid = existing.id if existing else None
        if conflict_policy == "skip":
            counts["would_skip"] += 1
            return "would_skip", eid
        counts["would_update"] += 1
        return "would_update", eid
    if kind == "ambiguous":
        counts["ambiguous"] += 1
        return "ambiguous", None
    counts["errors"] += 1
    return "error", None


def _tally_import(counts: dict[str, int], outcome: str) -> None:
    if outcome == "created":
        counts["created"] += 1
    elif outcome == "updated":
        counts["updated"] += 1
    elif outcome == "skipped":
        counts["skipped"] += 1
    elif outcome == "error":
        counts["errors"] += 1
    elif outcome == "review":
        counts["review"] += 1


async def _apply_match(
    db: AsyncSession,
    existing: Contact,
    mapped: Any,
    conflict_policy: str,
) -> tuple[str, str]:
    """Apply conflict_policy to an existing matched contact.

    Returns (outcome, message).
    """
    if conflict_policy == "skip":
        return "skipped", "conflict_policy=skip; no changes"

    _SCALAR_FIELDS = [
        "first_name", "last_name", "email", "phone", "gender",
        "birth_date", "contact_type", "contact_subtype", "nickname",
        "suffix", "street_address",
    ]

    changed = False

    for field in _SCALAR_FIELDS:
        val = getattr(mapped, field, None)
        if val is None or (isinstance(val, str) and not val.strip()):
            continue  # skip empty mapped values
        existing_val = getattr(existing, field, None)
        if conflict_policy == "fill":
            if existing_val is None or (isinstance(existing_val, str) and not existing_val.strip()):
                setattr(existing, field, val)
                changed = True
        else:  # update
            setattr(existing, field, val)
            changed = True

    # Merge custom_data
    mapped_cd: dict = getattr(mapped, "custom_data", None) or {}
    if mapped_cd:
        current_cd: dict = dict(existing.custom_data or {})
        if conflict_policy == "fill":
            for k, v in mapped_cd.items():
                if k not in current_cd or current_cd[k] is None:
                    current_cd[k] = v
                    changed = True
        else:  # update
            current_cd.update(mapped_cd)
            changed = True
        existing.custom_data = current_cd

    if not changed:
        return "skipped", "no changes"

    await db.flush()
    return "updated", f"updated contact id={existing.id}"


async def _create_contact_direct(
    db: AsyncSession,
    mapped: Any,
) -> tuple[str, Optional[int], str]:
    """INSERT a brand-new Contact row, bypassing any external_id lookup.

    Used when match_key is 'email' or 'name' and disposition returned 'new':
    calling upsert_contact in that path would re-match on external_id and could
    silently UPDATE an unrelated existing contact whose external_id happens to
    collide with the incoming row's external_id (security fix — S10 SEC-3).

    If the row's external_id is already owned by a different contact (UNIQUE
    violation), we insert without it rather than failing or mutating the owner.
    """
    from sqlalchemy.exc import IntegrityError  # noqa: PLC0415

    _SCALAR_FIELDS = [
        "first_name", "last_name", "email", "phone", "gender",
        "birth_date", "contact_type", "contact_subtype", "nickname",
        "suffix", "street_address",
    ]

    raw_ext = getattr(mapped, "external_id", None)
    ext_id: Optional[int] = None
    if raw_ext:
        try:
            ext_id = int(raw_ext)
        except (TypeError, ValueError):
            pass

    scalars = {
        f: getattr(mapped, f, None)
        for f in _SCALAR_FIELDS
        if getattr(mapped, f, None) is not None
    }
    custom_data = dict(getattr(mapped, "custom_data", None) or {})

    # Try inserting WITH external_id first (savepoint so we can rollback cleanly).
    if ext_id is not None:
        sp = await db.begin_nested()
        try:
            contact = Contact(
                external_id=ext_id,
                custom_data=custom_data or None,
                **scalars,
            )
            db.add(contact)
            await db.flush()
            await sp.commit()
            await db.refresh(contact)
            return "created", contact.id, f"created (match_key != external_id; new contact id={contact.id})"
        except IntegrityError:
            await sp.rollback()
            # external_id belongs to a different contact — insert WITHOUT it.
            ext_id = None
            logger.info(
                "_create_contact_direct: external_id collision; inserting without external_id"
            )

    contact = Contact(
        external_id=None,
        custom_data=custom_data or None,
        **scalars,
    )
    db.add(contact)
    await db.flush()
    await db.refresh(contact)
    msg = (
        "created (match_key != external_id; external_id omitted — collision with existing contact)"
        if ext_id is None and raw_ext
        else f"created (match_key != external_id; new contact id={contact.id})"
    )
    return "created", contact.id, msg


async def _enqueue_name_review(
    raw_name: str,
    batch_id: int,
    db: AsyncSession,
) -> Optional[int]:
    """Enqueue an ambiguous name match for manual review; return queue row id."""
    from app.services.name_match import match_name  # noqa: PLC0415

    try:
        result = await match_name(
            raw_name=raw_name,
            db=db,
            source="import",
            payload={"batch_id": batch_id},
            auto_enqueue=True,
        )
        return result.get("review_queue_id")
    except Exception as exc:
        logger.warning("_enqueue_name_review failed for %r: %s", raw_name, exc)
        return None


async def _resolve_contact_ref(
    contact_ref: str,
    db: AsyncSession,
) -> Optional[int]:
    """Resolve a contact_ref (external_id or internal id) to a Contact.id."""
    if not contact_ref:
        return None
    # Try as external_id (integer)
    try:
        ext_id = int(contact_ref)
        result = await db.execute(
            select(Contact.id).where(Contact.external_id == ext_id)
        )
        row = result.scalar_one_or_none()
        if row is not None:
            return row
        # Fall through: might be an internal id
        result2 = await db.execute(
            select(Contact.id).where(Contact.id == ext_id)
        )
        return result2.scalar_one_or_none()
    except (TypeError, ValueError):
        pass
    # Try as email
    result = await db.execute(
        select(Contact.id).where(Contact.email == contact_ref)
    )
    return result.scalar_one_or_none()


async def _resolve_event_ref(
    event_ref: str,
    db: AsyncSession,
) -> Optional[int]:
    """Resolve an event_ref (external_id string) to an Event.id."""
    if not event_ref:
        return None
    from app.models import Event
    try:
        ext_id = int(event_ref)
        result = await db.execute(
            select(Event.id).where(Event.external_id == ext_id)
        )
        row = result.scalar_one_or_none()
        if row is not None:
            return row
        # Try internal id
        result2 = await db.execute(
            select(Event.id).where(Event.id == ext_id)
        )
        return result2.scalar_one_or_none()
    except (TypeError, ValueError):
        return None


async def _map_row(
    entity: str,
    raw_row: dict,
    s06_map: dict,
) -> Any:
    """Map a raw row using the appropriate S06 mapper function."""
    if entity == "contact":
        return await map_contact_row(raw_row, s06_map, {})
    elif entity == "participant":
        return await map_participant_row(raw_row, s06_map, {})
    else:
        return MappingError(
            row_index=0,
            field="__row__",
            message=f"Unsupported entity for wizard import: {entity!r}",
            raw_payload="",
        )


async def _row_generator(
    fmt: str,
    staging_abs: Path,
    sheet: Optional[str],
    options: Optional[dict],
) -> AsyncIterator[dict]:
    """Yield raw row dicts from a staged file.

    UnicodeDecodeError on any mid-stream row raises ValueError so the caller
    can record an error outcome rather than propagating an unhandled 500.
    """
    if fmt == "csv":
        encoding = (options or {}).get("encoding", "utf-8-sig")
        delimiter = (options or {}).get("delimiter", ",")
        try:
            for row in iter_csv_rows(str(staging_abs), encoding, delimiter):
                yield row
        except UnicodeDecodeError as exc:
            raise ValueError(
                f"CSV decode error at byte offset {exc.start}: {exc.reason}. "
                "Re-save the file as UTF-8 and re-upload."
            ) from exc
    else:
        rows = await asyncio.to_thread(_iter_xlsx_sync, str(staging_abs), sheet)
        for row in rows:
            yield row
